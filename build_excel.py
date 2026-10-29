"""
Build Olist Excel Workbook (Olist_Ecommerce_Workbook.xlsx)
===========================================================
Author: Ayushi Singh, Data Analyst
Requirements:
- Sheets: Cover, Monthly, Category, State, Payment, Delivery_Review, RFM, Review_Dist, Cohort, Executive, Delivery_Dashboard.
- Cover states: source, delivered-only scope, 4 headline totals, Nov 2016 & Sep-Oct 2018 gaps, and fact_orders.csv notice.
- Executive sheet: Live formulas pointing to Monthly sheet (=SUM(Monthly!C2:C27) for GMV and =SUM(Monthly!B2:B27) for Orders).
- Delivery_Dashboard sheet: Live bar chart linked to Delivery_Review cells, plus headline "late 2.57 vs on-time 4.30".
- Native openpyxl charts linked to cells (never pasted images).
- Navy (#0F2040) headers, bold white text, freeze panes, autofit columns, and formatting.
- Verifies Gate 3 logic.
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, BarChart, Reference, Series
import pandas as pd

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CLEAN_DIR = os.path.join(BASE_DIR, "data", "clean")
WORKBOOK_PATH = os.path.join(BASE_DIR, "Olist_Ecommerce_Workbook.xlsx")

# Palette
NAVY_HEADER = "0F2040"
WHITE = "FFFFFF"
LIGHT_GRAY_FILL = "F8FAFC"
ACCENT_BLUE = "1E40AF"
SLATE_BLUE = "1E40AF"
CORAL_ACCENT = "E11D48"
BORDER_GRAY = "CBD5E1"
CARD_BG = "F1F5F9"

header_font = Font(name="Calibri", size=11, bold=True, color=WHITE)
header_fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")

card_title_font = Font(name="Calibri", size=10, bold=True, color="64748B")
card_val_font = Font(name="Calibri", size=18, bold=True, color=NAVY_HEADER)
card_fill = PatternFill(start_color=CARD_BG, end_color=CARD_BG, fill_type="solid")

thin_border = Border(
    left=Side(style='thin', color=BORDER_GRAY),
    right=Side(style='thin', color=BORDER_GRAY),
    top=Side(style='thin', color=BORDER_GRAY),
    bottom=Side(style='thin', color=BORDER_GRAY)
)


def autofit_and_style(ws, freeze_cell="A2"):
    ws.views.sheetView[0].showGridLines = True
    if freeze_cell:
        ws.freeze_panes = freeze_cell
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if cell.number_format and ('$' in cell.number_format or '%' in cell.number_format):
                max_len = max(max_len, len(val_str) + 4)
            else:
                max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)


def style_header_row(ws, row_idx=1, max_col=10):
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row_idx, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[row_idx].height = 28


def create_workbook():
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # ----------------------------------------------------
    # 1. COVER SHEET
    # ----------------------------------------------------
    ws_cover = wb.create_sheet(title="Cover")
    ws_cover.views.sheetView[0].showGridLines = True
    
    ws_cover.column_dimensions["B"].width = 28
    ws_cover.column_dimensions["C"].width = 65

    ws_cover["B2"] = "OLIST BRAZILIAN E-COMMERCE ANALYTICS WORKBOOK"
    ws_cover["B2"].font = Font(name="Calibri", size=16, bold=True, color=NAVY_HEADER)
    
    ws_cover["B3"] = "Historical E-Commerce Operations, RFM Customer Segments & Logistics Performance"
    ws_cover["B3"].font = Font(name="Calibri", size=11, italic=True, color="475569")

    metadata = [
        ("Data Source", "Kaggle: olistbr/brazilian-ecommerce (CC BY-NC-SA 4.0)"),
        ("Prepared By", "Ayushi Singh, Data Analyst"),
        ("Analysis Scope", "Delivered Orders Only (revenue/GMV includes delivered status only; shipped/cancelled/etc. are excluded from GMV)"),
        ("Delivered GMV", "R$ 13,221,498.11 (Sum of item prices for delivered orders)"),
        ("Delivered Freight", "R$ 2,202,965.54 (Reported separately, not conflated with GMV)"),
        ("Delivered Payments", "R$ 15,422,461.77 (Reconciliation difference: R$ 2,001.88)"),
        ("Delivered Orders", "96,478 delivered orders (93,358 unique customers)"),
        ("Average Order Value", "R$ 137.04 | Average Review Score: 4.16 / 5.00 | Late Rate: 8.11%"),
        ("Calendar Gaps Note", "Nov 2016 has 0 orders (known raw data gap); Sep-Oct 2018 has 0 delivered orders due to extraction cutoff."),
        ("Fact Table Notice", "The complete 96,478-row fact table is saved in data/clean/fact_orders.csv (and 110,197 items in fact_order_items.csv) to ensure workbook performance and usability.")
    ]

    ws_cover["B5"] = "Project Metadata & Scope Specification"
    ws_cover["B5"].font = Font(name="Calibri", size=12, bold=True, color=NAVY_HEADER)
    ws_cover.merge_cells("B5:C5")
    ws_cover["B5"].fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

    row_cur = 6
    for k, v in metadata:
        ws_cover.cell(row=row_cur, column=2, value=k).font = Font(name="Calibri", size=10, bold=True, color=NAVY_HEADER)
        ws_cover.cell(row=row_cur, column=2).fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        ws_cover.cell(row=row_cur, column=2).border = thin_border
        
        c = ws_cover.cell(row=row_cur, column=3, value=v)
        c.font = Font(name="Calibri", size=10, color="1E293B")
        c.border = thin_border
        c.alignment = Alignment(wrap_text=True)
        ws_cover.row_dimensions[row_cur].height = 24
        row_cur += 1

    # Table of Sheets
    row_cur += 2
    ws_cover.cell(row=row_cur, column=2, value="Workbook Structure & Navigation").font = Font(name="Calibri", size=12, bold=True, color=NAVY_HEADER)
    ws_cover.merge_cells(f"B{row_cur}:C{row_cur}")
    ws_cover.cell(row=row_cur, column=2).fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    row_cur += 1

    sheets_desc = [
        ("Executive", "Executive KPI summary dashboard with live formulas linked to Monthly sheet and native Excel charts."),
        ("Delivery_Dashboard", "Logistics & customer satisfaction dashboard showing review scores by delivery bucket."),
        ("Monthly", "Monthly delivered orders, GMV, freight, AOV, late rates, and average reviews (Sep 2016 – Oct 2018)."),
        ("Category", "Performance of 72 product categories by items sold, GMV, freight, and average item price."),
        ("State", "Customer geography breakdown across 27 Brazilian states with GMV, orders, and late rates."),
        ("Payment", "Payment methods breakdown: credit card, boleto, voucher, debit card with installments."),
        ("Delivery_Review", "Customer review score and 5-star share segmented across 6 delivery duration buckets."),
        ("RFM", "RFM customer segments breakdown (Champions, Loyal, Recent/Potential, Regular, Lost/Hibernating)."),
        ("Review_Dist", "Order review score distribution (1 to 5 stars) with counts and percentages."),
        ("Cohort", "Customer cohort retention matrix (first purchase month M0 through M20).")
    ]

    for s_name, s_desc in sheets_desc:
        ws_cover.cell(row=row_cur, column=2, value=s_name).font = Font(name="Calibri", size=10, bold=True, color=SLATE_BLUE)
        ws_cover.cell(row=row_cur, column=2).border = thin_border
        c = ws_cover.cell(row=row_cur, column=3, value=s_desc)
        c.font = Font(name="Calibri", size=10, color="1E293B")
        c.border = thin_border
        ws_cover.row_dimensions[row_cur].height = 20
        row_cur += 1

    # ----------------------------------------------------
    # 2. MONTHLY SHEET
    # ----------------------------------------------------
    ws_monthly = wb.create_sheet(title="Monthly")
    df_monthly = pd.read_csv(os.path.join(CLEAN_DIR, "monthly.csv"))
    
    headers_monthly = ["Purchase Month", "Delivered Orders", "Delivered GMV (R$)", "Delivered Freight (R$)", "Avg Order Value (R$)", "Late Orders", "Late Rate (%)", "Avg Delivery Days", "Avg Review Score"]
    for col_idx, h in enumerate(headers_monthly, 1):
        ws_monthly.cell(row=1, column=col_idx, value=h)
    style_header_row(ws_monthly, 1, len(headers_monthly))

    for r_idx, row in df_monthly.iterrows():
        r = r_idx + 2
        ws_monthly.cell(row=r, column=1, value=str(row["purchase_month"])).alignment = Alignment(horizontal="center")
        c2 = ws_monthly.cell(row=r, column=2, value=int(row["orders"]))
        c2.number_format = "#,##0"
        c3 = ws_monthly.cell(row=r, column=3, value=float(row["gmv"]))
        c3.number_format = "$#,##0.00"
        c4 = ws_monthly.cell(row=r, column=4, value=float(row["freight"]))
        c4.number_format = "$#,##0.00"
        c5 = ws_monthly.cell(row=r, column=5, value=float(row["avg_order_value"]))
        c5.number_format = "$#,##0.00"
        c6 = ws_monthly.cell(row=r, column=6, value=int(row["late_orders"]))
        c6.number_format = "#,##0"
        c7 = ws_monthly.cell(row=r, column=7, value=float(row["late_rate"])/100.0)
        c7.number_format = "0.00%"
        c8 = ws_monthly.cell(row=r, column=8, value=float(row["avg_delivery_days"]))
        c8.number_format = "0.00"
        c9 = ws_monthly.cell(row=r, column=9, value=float(row["avg_review"]))
        c9.number_format = "0.00"
        for col_idx in range(1, len(headers_monthly) + 1):
            ws_monthly.cell(row=r, column=col_idx).border = thin_border
            
    autofit_and_style(ws_monthly, "A2")

    # ----------------------------------------------------
    # 3. CATEGORY SHEET
    # ----------------------------------------------------
    ws_cat = wb.create_sheet(title="Category")
    df_cat = pd.read_csv(os.path.join(CLEAN_DIR, "category.csv"))
    headers_cat = ["English Category", "Items Sold", "Delivered GMV (R$)", "Delivered Freight (R$)", "Avg Item Price (R$)", "Share of GMV (%)"]
    for col_idx, h in enumerate(headers_cat, 1):
        ws_cat.cell(row=1, column=col_idx, value=h)
    style_header_row(ws_cat, 1, len(headers_cat))

    for r_idx, row in df_cat.iterrows():
        r = r_idx + 2
        ws_cat.cell(row=r, column=1, value=str(row["category_en"]))
        ws_cat.cell(row=r, column=2, value=int(row["items_sold"])).number_format = "#,##0"
        ws_cat.cell(row=r, column=3, value=float(row["gmv"])).number_format = "$#,##0.00"
        ws_cat.cell(row=r, column=4, value=float(row["freight"])).number_format = "$#,##0.00"
        ws_cat.cell(row=r, column=5, value=float(row["avg_item_price"])).number_format = "$#,##0.00"
        ws_cat.cell(row=r, column=6, value=float(row["share_of_gmv"])/100.0).number_format = "0.00%"
        for col_idx in range(1, len(headers_cat) + 1):
            ws_cat.cell(row=r, column=col_idx).border = thin_border
    autofit_and_style(ws_cat, "A2")

    # ----------------------------------------------------
    # 4. STATE SHEET
    # ----------------------------------------------------
    ws_state = wb.create_sheet(title="State")
    df_state = pd.read_csv(os.path.join(CLEAN_DIR, "state.csv"))
    headers_state = ["Customer State", "Delivered Orders", "Delivered GMV (R$)", "Delivered Freight (R$)", "Late Orders", "Late Rate (%)", "Avg Delivery Days", "Avg Review Score", "Share of GMV (%)"]
    for col_idx, h in enumerate(headers_state, 1):
        ws_state.cell(row=1, column=col_idx, value=h)
    style_header_row(ws_state, 1, len(headers_state))

    for r_idx, row in df_state.iterrows():
        r = r_idx + 2
        ws_state.cell(row=r, column=1, value=str(row["customer_state"])).alignment = Alignment(horizontal="center")
        ws_state.cell(row=r, column=2, value=int(row["orders"])).number_format = "#,##0"
        ws_state.cell(row=r, column=3, value=float(row["gmv"])).number_format = "$#,##0.00"
        ws_state.cell(row=r, column=4, value=float(row["freight"])).number_format = "$#,##0.00"
        ws_state.cell(row=r, column=5, value=int(row["late_orders"])).number_format = "#,##0"
        ws_state.cell(row=r, column=6, value=float(row["late_rate"])/100.0).number_format = "0.00%"
        ws_state.cell(row=r, column=7, value=float(row["avg_delivery_days"])).number_format = "0.00"
        ws_state.cell(row=r, column=8, value=float(row["avg_review"])).number_format = "0.00"
        ws_state.cell(row=r, column=9, value=float(row["share_of_gmv"])/100.0).number_format = "0.00%"
        for col_idx in range(1, len(headers_state) + 1):
            ws_state.cell(row=r, column=col_idx).border = thin_border
    autofit_and_style(ws_state, "A2")

    # ----------------------------------------------------
    # 5. PAYMENT SHEET
    # ----------------------------------------------------
    ws_pay = wb.create_sheet(title="Payment")
    df_pay = pd.read_csv(os.path.join(CLEAN_DIR, "payment_type.csv"))
    headers_pay = ["Payment Type", "Total Value (R$)", "Avg Installments", "Payment Count", "Share of Total (%)"]
    for col_idx, h in enumerate(headers_pay, 1):
        ws_pay.cell(row=1, column=col_idx, value=h)
    style_header_row(ws_pay, 1, len(headers_pay))

    for r_idx, row in df_pay.iterrows():
        r = r_idx + 2
        ws_pay.cell(row=r, column=1, value=str(row["payment_type"]))
        ws_pay.cell(row=r, column=2, value=float(row["total_value"])).number_format = "$#,##0.00"
        ws_pay.cell(row=r, column=3, value=float(row["avg_installments"])).number_format = "0.00"
        ws_pay.cell(row=r, column=4, value=int(row["payment_count"])).number_format = "#,##0"
        ws_pay.cell(row=r, column=5, value=float(row["share_of_total"])/100.0).number_format = "0.00%"
        for col_idx in range(1, len(headers_pay) + 1):
            ws_pay.cell(row=r, column=col_idx).border = thin_border
    autofit_and_style(ws_pay, "A2")

    # ----------------------------------------------------
    # 6. DELIVERY_REVIEW SHEET
    # ----------------------------------------------------
    ws_deliv_rev = wb.create_sheet(title="Delivery_Review")
    df_deliv_rev = pd.read_csv(os.path.join(CLEAN_DIR, "delivery_vs_review.csv"))
    headers_deliv_rev = ["Delivery Bucket", "Delivered Orders", "Avg Review Score", "5-Star Orders Count", "5-Star Share (%)", "Share of Orders (%)"]
    for col_idx, h in enumerate(headers_deliv_rev, 1):
        ws_deliv_rev.cell(row=1, column=col_idx, value=h)
    style_header_row(ws_deliv_rev, 1, len(headers_deliv_rev))

    for r_idx, row in df_deliv_rev.iterrows():
        r = r_idx + 2
        ws_deliv_rev.cell(row=r, column=1, value=str(row["delivery_bucket"])).alignment = Alignment(horizontal="center")
        ws_deliv_rev.cell(row=r, column=2, value=int(row["orders"])).number_format = "#,##0"
        ws_deliv_rev.cell(row=r, column=3, value=float(row["avg_review"])).number_format = "0.00"
        ws_deliv_rev.cell(row=r, column=4, value=int(row["five_star_count"])).number_format = "#,##0"
        ws_deliv_rev.cell(row=r, column=5, value=float(row["pct_5star"])/100.0).number_format = "0.00%"
        ws_deliv_rev.cell(row=r, column=6, value=float(row["share_of_orders"])/100.0).number_format = "0.00%"
        for col_idx in range(1, len(headers_deliv_rev) + 1):
            ws_deliv_rev.cell(row=r, column=col_idx).border = thin_border
    autofit_and_style(ws_deliv_rev, "A2")

    # ----------------------------------------------------
    # 7. RFM SHEET
    # ----------------------------------------------------
    ws_rfm = wb.create_sheet(title="RFM")
    df_rfm = pd.read_csv(os.path.join(CLEAN_DIR, "rfm_segments.csv"))
    headers_rfm = ["Segment", "Customers Count", "Total GMV (R$)", "Avg Spend (R$)", "Share of Customers (%)", "Share of GMV (%)"]
    for col_idx, h in enumerate(headers_rfm, 1):
        ws_rfm.cell(row=1, column=col_idx, value=h)
    style_header_row(ws_rfm, 1, len(headers_rfm))

    for r_idx, row in df_rfm.iterrows():
        r = r_idx + 2
        ws_rfm.cell(row=r, column=1, value=str(row["segment"]))
        ws_rfm.cell(row=r, column=2, value=int(row["customers"])).number_format = "#,##0"
        ws_rfm.cell(row=r, column=3, value=float(row["total_gmv"])).number_format = "$#,##0.00"
        ws_rfm.cell(row=r, column=4, value=float(row["avg_spend"])).number_format = "$#,##0.00"
        ws_rfm.cell(row=r, column=5, value=float(row["share_of_customers"])/100.0).number_format = "0.00%"
        ws_rfm.cell(row=r, column=6, value=float(row["share_of_gmv"])/100.0).number_format = "0.00%"
        for col_idx in range(1, len(headers_rfm) + 1):
            ws_rfm.cell(row=r, column=col_idx).border = thin_border
    autofit_and_style(ws_rfm, "A2")

    # ----------------------------------------------------
    # 8. REVIEW_DIST SHEET
    # ----------------------------------------------------
    ws_rev_dist = wb.create_sheet(title="Review_Dist")
    df_rev_dist = pd.read_csv(os.path.join(CLEAN_DIR, "review_distribution.csv"))
    headers_rev_dist = ["Review Score", "Delivered Orders Count", "Share of Reviews (%)", "Share of All Delivered Orders (%)"]
    for col_idx, h in enumerate(headers_rev_dist, 1):
        ws_rev_dist.cell(row=1, column=col_idx, value=h)
    style_header_row(ws_rev_dist, 1, len(headers_rev_dist))

    for r_idx, row in df_rev_dist.iterrows():
        r = r_idx + 2
        ws_rev_dist.cell(row=r, column=1, value=int(row["review_score"])).alignment = Alignment(horizontal="center")
        ws_rev_dist.cell(row=r, column=2, value=int(row["count"])).number_format = "#,##0"
        ws_rev_dist.cell(row=r, column=3, value=float(row["pct_of_reviews"])/100.0).number_format = "0.00%"
        ws_rev_dist.cell(row=r, column=4, value=float(row["pct_of_delivered_orders"])/100.0).number_format = "0.00%"
        for col_idx in range(1, len(headers_rev_dist) + 1):
            ws_rev_dist.cell(row=r, column=col_idx).border = thin_border
    autofit_and_style(ws_rev_dist, "A2")

    # ----------------------------------------------------
    # 9. COHORT SHEET
    # ----------------------------------------------------
    ws_cohort = wb.create_sheet(title="Cohort")
    df_cohort = pd.read_csv(os.path.join(CLEAN_DIR, "cohort_retention.csv"))
    for col_idx, col_name in enumerate(df_cohort.columns, 1):
        ws_cohort.cell(row=1, column=col_idx, value=f"M+{col_name}" if col_name.isdigit() else str(col_name))
    style_header_row(ws_cohort, 1, len(df_cohort.columns))

    for r_idx, row in df_cohort.iterrows():
        r = r_idx + 2
        ws_cohort.cell(row=r, column=1, value=str(row["cohort_month"]))
        ws_cohort.cell(row=r, column=2, value=int(row["cohort_size"])).number_format = "#,##0"
        for col_idx, c_name in enumerate(df_cohort.columns[2:], 3):
            val = row[c_name]
            cell = ws_cohort.cell(row=r, column=col_idx)
            if pd.notna(val):
                cell.value = int(val)
                cell.number_format = "#,##0"
            cell.border = thin_border
    autofit_and_style(ws_cohort, "A2")

    # ----------------------------------------------------
    # 10. EXECUTIVE DASHBOARD SHEET
    # ----------------------------------------------------
    ws_exec = wb.create_sheet(title="Executive")
    ws_exec.views.sheetView[0].showGridLines = True
    ws_exec.column_dimensions["A"].width = 4
    ws_exec.column_dimensions["B"].width = 24
    ws_exec.column_dimensions["C"].width = 24
    ws_exec.column_dimensions["D"].width = 24
    ws_exec.column_dimensions["E"].width = 24

    ws_exec["B2"] = "EXECUTIVE E-COMMERCE PERFORMANCE DASHBOARD"
    ws_exec["B2"].font = Font(name="Calibri", size=16, bold=True, color=NAVY_HEADER)
    ws_exec["B3"] = "Real Kaggle Olist Dataset (Delivered Orders Only | Sep 2016 – Aug 2018)"
    ws_exec["B3"].font = Font(name="Calibri", size=10, italic=True, color="475569")

    # KPI Cards Row
    # Card 1: Delivered GMV (LIVE FORMULA pointing to Monthly sheet)
    ws_exec["B5"] = "DELIVERED GMV"
    ws_exec["B5"].font = card_title_font
    ws_exec["B5"].fill = card_fill
    ws_exec["B5"].border = thin_border
    ws_exec["B5"].alignment = Alignment(horizontal="center")
    
    ws_exec["B6"] = "=SUM(Monthly!C2:C27)"
    ws_exec["B6"].font = card_val_font
    ws_exec["B6"].fill = card_fill
    ws_exec["B6"].border = thin_border
    ws_exec["B6"].number_format = "$#,##0.00"
    ws_exec["B6"].alignment = Alignment(horizontal="center")

    # Card 2: Delivered Orders (LIVE FORMULA pointing to Monthly sheet)
    ws_exec["C5"] = "DELIVERED ORDERS"
    ws_exec["C5"].font = card_title_font
    ws_exec["C5"].fill = card_fill
    ws_exec["C5"].border = thin_border
    ws_exec["C5"].alignment = Alignment(horizontal="center")
    
    ws_exec["C6"] = "=SUM(Monthly!B2:B27)"
    ws_exec["C6"].font = card_val_font
    ws_exec["C6"].fill = card_fill
    ws_exec["C6"].border = thin_border
    ws_exec["C6"].number_format = "#,##0"
    ws_exec["C6"].alignment = Alignment(horizontal="center")

    # Card 3: Avg Order Value (LIVE FORMULA)
    ws_exec["D5"] = "AVG ORDER VALUE (AOV)"
    ws_exec["D5"].font = card_title_font
    ws_exec["D5"].fill = card_fill
    ws_exec["D5"].border = thin_border
    ws_exec["D5"].alignment = Alignment(horizontal="center")
    
    ws_exec["D6"] = "=B6/C6"
    ws_exec["D6"].font = card_val_font
    ws_exec["D6"].fill = card_fill
    ws_exec["D6"].border = thin_border
    ws_exec["D6"].number_format = "$#,##0.00"
    ws_exec["D6"].alignment = Alignment(horizontal="center")

    # Card 4: Avg Review Score (Order-weighted total from summary)
    ws_exec["E5"] = "AVG REVIEW SCORE"
    ws_exec["E5"].font = card_title_font
    ws_exec["E5"].fill = card_fill
    ws_exec["E5"].border = thin_border
    ws_exec["E5"].alignment = Alignment(horizontal="center")
    
    ws_exec["E6"] = 4.16
    ws_exec["E6"].font = card_val_font
    ws_exec["E6"].fill = card_fill
    ws_exec["E6"].border = thin_border
    ws_exec["E6"].number_format = "0.00"
    ws_exec["E6"].alignment = Alignment(horizontal="center")

    # Verification Note
    ws_exec["B8"] = "Gate 2 Reconciliation Note: Delivered Payments = R$ 15,422,461.77 | Freight = R$ 2,202,965.54 | Discrepancy = R$ 2,001.88"
    ws_exec["B8"].font = Font(name="Calibri", size=9, bold=True, color="047857")

    # Native Excel Chart 1: Monthly GMV Line Chart (Linked to Monthly sheet)
    chart_line = LineChart()
    chart_line.title = "Monthly Delivered GMV (R$)"
    chart_line.style = 13
    chart_line.y_axis.title = "Delivered GMV (R$)"
    chart_line.x_axis.title = "Purchase Month"
    chart_line.width = 18
    chart_line.height = 10

    # Data: Monthly GMV (Column C, rows 1 to 25 -> active months Sep 2016 to Aug 2018)
    data_ref = Reference(ws_monthly, min_col=3, min_row=1, max_row=25)
    cats_ref = Reference(ws_monthly, min_col=1, min_row=2, max_row=25)
    chart_line.add_data(data_ref, titles_from_data=True)
    chart_line.set_categories(cats_ref)
    ws_exec.add_chart(chart_line, "B10")

    # Native Excel Chart 2: Top 10 Categories Bar Chart (Linked to Category sheet)
    chart_bar = BarChart()
    chart_bar.type = "bar"
    chart_bar.title = "Top 10 Categories by GMV (R$)"
    chart_bar.style = 10
    chart_bar.y_axis.title = "Category"
    chart_bar.x_axis.title = "Delivered GMV (R$)"
    chart_bar.width = 18
    chart_bar.height = 10

    data_cat = Reference(ws_cat, min_col=3, min_row=1, max_row=11)
    cats_cat = Reference(ws_cat, min_col=1, min_row=2, max_row=11)
    chart_bar.add_data(data_cat, titles_from_data=True)
    chart_bar.set_categories(cats_cat)
    ws_exec.add_chart(chart_bar, "B31")

    # ----------------------------------------------------
    # 11. DELIVERY_DASHBOARD SHEET
    # ----------------------------------------------------
    ws_deliv_dash = wb.create_sheet(title="Delivery_Dashboard")
    ws_deliv_dash.views.sheetView[0].showGridLines = True
    ws_deliv_dash.column_dimensions["A"].width = 4
    ws_deliv_dash.column_dimensions["B"].width = 24
    ws_deliv_dash.column_dimensions["C"].width = 24
    ws_deliv_dash.column_dimensions["D"].width = 24
    ws_deliv_dash.column_dimensions["E"].width = 24

    ws_deliv_dash["B2"] = "DELIVERY LOGISTICS & SATISFACTION DASHBOARD"
    ws_deliv_dash["B2"].font = Font(name="Calibri", size=16, bold=True, color=NAVY_HEADER)
    
    ws_deliv_dash["B3"] = "Headline Satisfaction Gap: Late Orders 2.57 stars vs On-Time Orders 4.30 stars (Late Rate: 8.11%)"
    ws_deliv_dash["B3"].font = Font(name="Calibri", size=11, bold=True, color=CORAL_ACCENT)

    # Mini KPI Cards
    dash_kpis = [
        ("B5", "ON-TIME ORDERS AVG REVIEW", "4.30 / 5.00"),
        ("C5", "LATE ORDERS AVG REVIEW", "2.57 / 5.00"),
        ("D5", "NATIONAL LATE RATE", "8.11%"),
        ("E5", "MEDIAN DELIVERY TIME", "10.22 Days")
    ]
    for cell_pos, title, val in dash_kpis:
        col_letter = cell_pos[0]
        ws_deliv_dash[f"{col_letter}5"] = title
        ws_deliv_dash[f"{col_letter}5"].font = card_title_font
        ws_deliv_dash[f"{col_letter}5"].fill = card_fill
        ws_deliv_dash[f"{col_letter}5"].border = thin_border
        ws_deliv_dash[f"{col_letter}5"].alignment = Alignment(horizontal="center")

        ws_deliv_dash[f"{col_letter}6"] = val
        ws_deliv_dash[f"{col_letter}6"].font = card_val_font
        ws_deliv_dash[f"{col_letter}6"].fill = card_fill
        ws_deliv_dash[f"{col_letter}6"].border = thin_border
        ws_deliv_dash[f"{col_letter}6"].alignment = Alignment(horizontal="center")

    # Native Excel Chart 3: Avg Review by Delivery Bucket (Linked to Delivery_Review sheet)
    chart_deliv = BarChart()
    chart_deliv.type = "col"
    chart_deliv.title = "Average Customer Review Score by Delivery Bucket"
    chart_deliv.style = 11
    chart_deliv.y_axis.title = "Average Review Score (1-5)"
    chart_deliv.x_axis.title = "Delivery Duration Bucket"
    chart_deliv.width = 18
    chart_deliv.height = 11

    data_deliv = Reference(ws_deliv_rev, min_col=3, min_row=1, max_row=7)
    cats_deliv = Reference(ws_deliv_rev, min_col=1, min_row=2, max_row=7)
    chart_deliv.add_data(data_deliv, titles_from_data=True)
    chart_deliv.set_categories(cats_deliv)
    ws_deliv_dash.add_chart(chart_deliv, "B9")

    # Save workbook
    wb.save(WORKBOOK_PATH)
    print(f"Workbook successfully saved to {WORKBOOK_PATH}")


def verify_gate_3():
    """
    Verification Gate 3:
    Open the workbook logic in Python (openpyxl, data_only=False),
    confirm the Executive formulas point at the Monthly sheet,
    and confirm the Monthly sums equal Gate 2 metrics.
    """
    print("=" * 60)
    print("VERIFICATION GATE 3 — EXCEL WORKBOOK LOGIC & FORMULA CHECKS")
    print("=" * 60)
    
    wb_formulas = openpyxl.load_workbook(WORKBOOK_PATH, data_only=False)
    ws_exec = wb_formulas["Executive"]
    
    # Check formula strings
    gmv_formula = ws_exec["B6"].value
    orders_formula = ws_exec["C6"].value
    aov_formula = ws_exec["D6"].value

    print(f"Executive B6 formula: {gmv_formula}")
    print(f"Executive C6 formula: {orders_formula}")
    print(f"Executive D6 formula: {aov_formula}")

    assert "Monthly!" in str(gmv_formula), f"GMV formula does not point to Monthly sheet: {gmv_formula}"
    assert "Monthly!" in str(orders_formula), f"Orders formula does not point to Monthly sheet: {orders_formula}"

    # Evaluate Monthly sheet values
    ws_monthly = wb_formulas["Monthly"]
    monthly_orders_sum = sum(ws_monthly.cell(row=r, column=2).value for r in range(2, 28))
    monthly_gmv_sum = sum(ws_monthly.cell(row=r, column=3).value for r in range(2, 28))

    print(f"Monthly Orders Sum: {monthly_orders_sum:,} (Expected: 96,478)")
    print(f"Monthly GMV Sum:    R$ {monthly_gmv_sum:,.2f} (Expected: 13,221,498.11)")

    assert monthly_orders_sum == 96478, f"Monthly orders sum mismatch: {monthly_orders_sum}"
    assert round(monthly_gmv_sum, 2) == 13221498.11, f"Monthly GMV sum mismatch: {monthly_gmv_sum}"

    print("Gate 3 Verification: ALL CHECKS PASSED (Formulas valid, Monthly sums match Gate 2)!\n")


if __name__ == "__main__":
    create_workbook()
    verify_gate_3()
